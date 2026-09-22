"""Download every URL the curated sheet cites and pull the titles out of it.

Written for step 3 of the review: the contribution and event names in the sheet
came out of one free-text column typed by hand over four years, so they carry
typos, truncations and house spellings. The pages already cited in the sheet
are the authority for what a talk was actually called, and this script fetches
them once into a local cache so the comparison can be run and re-run offline.

Output: local/source_titles.json, mapping each URL to the candidate titles
found on that page. Nothing in the sheet is changed here.

    python scripts/fetch_source_titles.py [--refresh]
"""

from __future__ import annotations

import concurrent.futures
import hashlib
import html
import json
import re
import subprocess
import sys
import warnings
from pathlib import Path
from urllib.parse import urlparse

warnings.filterwarnings('ignore', module='openpyxl')

import openpyxl  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
LOCAL = ROOT / 'local'
CACHE = LOCAL / '.page_cache'
OUT = LOCAL / 'source_titles.json'
SHEET = 'Events'

USER_AGENT = (
    'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) '
    'Chrome/126.0 Safari/537.36'
)
TIMEOUT = 25
WORKERS = 8
# A candidate title shorter than this is a label, longer than this is a page.
MIN_TITLE, MAX_TITLE = 4, 300
URL_RE = re.compile(r'https?://[^\s<>"\']+')


def newest_sheet() -> Path:
    candidates = sorted(LOCAL.glob('Events database - CURATED_*.xlsx'))
    if not candidates:
        raise SystemExit('no curated sheet found; run build_curated_sheet.py first')
    return candidates[-1]


def collect_urls(path: Path) -> dict[str, list[int]]:
    """URL -> the sheet rows citing it."""
    sheet = openpyxl.load_workbook(path, data_only=True)[SHEET]
    header = {sheet.cell(1, c).value: c for c in range(1, sheet.max_column + 1)}
    columns = [header['+ | Source URLs'], header['J | Link to relevant output 1']]
    found: dict[str, list[int]] = {}
    for row in range(2, sheet.max_row + 1):
        for column in columns:
            for url in URL_RE.findall(str(sheet.cell(row, column).value or '')):
                found.setdefault(url.rstrip('.,);'), []).append(row)
    return found


def cache_path(url: str) -> Path:
    return CACHE / f'{hashlib.sha1(url.encode()).hexdigest()}.html'


def fetch(url: str) -> str:
    """Return the page body, from the cache when it is already there."""
    target = cache_path(url)
    if target.exists():
        return target.read_text(encoding='utf-8', errors='replace')
    try:
        result = subprocess.run(
            ['curl', '-sL', '--compressed', '--max-time', str(TIMEOUT),
             '-A', USER_AGENT, url],
            capture_output=True, timeout=TIMEOUT + 10, check=False,
        )
        body = result.stdout.decode('utf-8', errors='replace')
    except (subprocess.TimeoutExpired, OSError):
        body = ''
    target.write_text(body, encoding='utf-8')
    return body


def text_of(fragment: str) -> str:
    fragment = re.sub(r'<[^>]+>', ' ', fragment)
    fragment = html.unescape(fragment)
    return re.sub(r'\s+', ' ', fragment).strip()


def titles_from(body: str, url: str) -> list[str]:
    """Candidate titles on the page, best guess first.

    Deliberately generous: the point is to offer a human every string on the
    page that could be the real name of the talk, not to pick one.
    """
    out: list[str] = []

    def add(value: str) -> None:
        value = text_of(value)
        if MIN_TITLE <= len(value) <= MAX_TITLE and value not in out:
            out.append(value)

    for pattern in (
        r'<meta[^>]+property=["\']og:title["\'][^>]+content=["\']([^"\']+)',
        r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+property=["\']og:title["\']',
        r'<meta[^>]+name=["\']citation_title["\'][^>]+content=["\']([^"\']+)',
        r'<meta[^>]+name=["\']twitter:title["\'][^>]+content=["\']([^"\']+)',
        r'<title[^>]*>(.*?)</title>',
    ):
        for hit in re.findall(pattern, body, re.I | re.S):
            add(hit)

    for level in ('h1', 'h2', 'h3', 'h4'):
        for hit in re.findall(rf'<{level}[^>]*>(.*?)</{level}>', body, re.I | re.S):
            add(hit)

    host = urlparse(url).netloc
    # Indico marks up each contribution; the talk we want is one of them.
    if 'indico' in body.lower() or host.startswith(('events.', 'indico.')):
        for hit in re.findall(
            r'class="[^"]*(?:contribution-title|timetable-title|item-title)[^"]*"[^>]*>(.*?)<',
            body, re.I | re.S,
        ):
            add(hit)
    # DPG programme pages put the talk title in the abstract block.
    if 'dpg-verhandlungen' in host:
        for hit in re.findall(r'<b>(.*?)</b>', body, re.I | re.S):
            add(hit)
    return out[:60]


def main() -> int:
    refresh = '--refresh' in sys.argv
    CACHE.mkdir(parents=True, exist_ok=True)
    if refresh:
        for stale in CACHE.glob('*.html'):
            stale.unlink()

    urls = collect_urls(newest_sheet())
    print(f'{len(urls)} distinct URLs cited by the sheet')

    bodies: dict[str, str] = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = {pool.submit(fetch, url): url for url in urls}
        for done, future in enumerate(concurrent.futures.as_completed(futures), 1):
            url = futures[future]
            bodies[url] = future.result()
            if done % 50 == 0:
                print(f'  fetched {done}/{len(urls)}')

    result = {}
    empty = 0
    for url, body in bodies.items():
        found = titles_from(body, url) if body else []
        if not found:
            empty += 1
        result[url] = {'rows': sorted(set(urls[url])), 'titles': found,
                       'bytes': len(body)}
    OUT.write_text(json.dumps(result, indent=1, ensure_ascii=False), encoding='utf-8')
    print(f'{len(result) - empty} pages yielded a title, {empty} did not')
    print(f'written to {OUT}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
